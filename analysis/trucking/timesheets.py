"""Consolidated foreman-timesheet loader.

Source: Trucking_History_Cumulative_Source_Data.xlsx / "Consolidated Data" — every
row from every weekly foreman timesheet export, already merged and tagged with the
invoice it coded. This is the audit trail behind the invoice-line cost codes, and
the basis for real cycle times (HR vs LOAD rows per truck-day) and the foreman
scorecard.
"""
import re
import pandas as pd
import openpyxl

COLS = ["SourceFile", "InvoiceSupported", "Coverage", "Date", "FieldLogID",
        "Foreman", "JobNumber", "Subcontractor", "TruckRaw", "Code",
        "CodeDesc", "Qty", "UM", "UnitCost", "TotalCost"]


def _truck(raw):
    if not raw:
        return ""
    tok = str(raw).strip().split()[0]           # "3C22 Three C's Prop" -> "3C22"
    return re.sub(r"-?Q$", "", tok, flags=re.I)  # strip quad-axle suffix


def _num(x):
    try:
        return float(str(x).replace("$", "").replace(",", ""))
    except Exception:
        return None


def load_timesheets(path: str):
    """Returns (DataFrame, log). Empty frame + status when the file is absent."""
    log = []
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except FileNotFoundError:
        return pd.DataFrame(columns=COLS), ["no timesheet file at %s" % path]
    if "Consolidated Data" not in wb.sheetnames:
        return pd.DataFrame(columns=COLS), ["'Consolidated Data' sheet not found"]
    ws = wb["Consolidated Data"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    recs = []
    for r in rows[1:]:
        d = dict(zip(header, r))
        recs.append({
            "SourceFile": d.get("Source File"),
            "InvoiceSupported": d.get("Invoice Supported"),
            "Coverage": d.get("Coverage Period"),
            "Date": pd.to_datetime(d.get("Date"), errors="coerce"),
            "FieldLogID": d.get("Field Log ID"),
            "Foreman": (d.get("Foreman") or "").strip(),
            "JobNumber": d.get("Job Number"),
            "Subcontractor": d.get("Trucking Subcontractor"),
            "TruckRaw": d.get("Truck ID Listing"),
            "Truck": _truck(d.get("Truck ID Listing")),
            "Code": str(d.get("Account ID") or "").strip().replace("-", "."),
            "CodeDesc": d.get("Account Description"),
            "Qty": _num(d.get("Quantity")),
            "UM": str(d.get("UM") or "").strip().upper(),
            "UnitCost": _num(d.get("Unit Cost")),
            "TotalCost": _num(d.get("Total Cost")),
        })
    df = pd.DataFrame(recs)
    log.append("timesheet rows loaded: %d" % len(df))
    log.append("unique trucks: %d, foremen: %d" %
               (df["Truck"].nunique(), df["Foreman"].nunique()))
    return df, log
